"""Document loaders: the lab corpus, private notes, and the paper library."""
from __future__ import annotations

import hashlib
import html
import json
import re
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass, field, replace
from html.parser import HTMLParser
from pathlib import Path
from typing import Iterator

import yaml

CACHE_DIR = Path.home() / ".cache" / "graphrag" / "pdf-text"
BLOCK_TAGS = {
    "p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6",
    "section", "article", "header", "footer", "pre", "blockquote",
}
SKIP_DIRS = {".git", ".hg", ".svn", "__pycache__", ".venv", "venv", "node_modules"}


# One normalized document ready for chunking and ingestion.
@dataclass
class SourceDocument:
    id: str
    corpus: str
    corpus_version: str
    title: str
    section: str
    path: str
    sha256: str
    source: dict
    text: str
    page_map: list[tuple[int, int]] | None = field(default=None)


# Strip HTML tags while keeping block boundaries and readable text.
class _TextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.skip = 0

    # Start collecting or skipping text at script/style boundaries.
    def handle_starttag(self, tag: str, attrs):
        if tag in ("script", "style"):
            self.skip += 1
        elif tag in BLOCK_TAGS:
            self.parts.append("\n")

    # Stop skipping at the matching closing tag.
    def handle_endtag(self, tag: str):
        if tag in ("script", "style") and self.skip:
            self.skip -= 1
        elif tag in BLOCK_TAGS:
            self.parts.append("\n")

    # Keep text nodes outside scripts and styles.
    def handle_data(self, data: str):
        if not self.skip:
            self.parts.append(data)

    # Return the accumulated plain text.
    def text(self) -> str:
        return "".join(self.parts)


# Collapse whitespace without losing paragraphs.
def normalize_text(text: str) -> str:
    # PostgreSQL text cannot store NUL bytes; some PDFs contain them.
    text = text.replace("\x00", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# Hash the text and the metadata that affects retrieval.
def content_hash(text: str, metadata: dict | None = None) -> str:
    payload = text + "\n" + json.dumps(metadata or {}, sort_keys=True, ensure_ascii=False)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


# Return the first Markdown heading, or a readable file name.
def guess_title(path: Path, text: str) -> str:
    for line in text.splitlines()[:20]:
        if line.startswith("# "):
            return line[2:].strip()[:200]
    return path.stem.replace("_", " ").replace("-", " ").strip()[:200]


# Strip script/style blocks and tags from an HTML file.
def html_to_text(raw: str) -> str:
    parser = _TextExtractor()
    parser.feed(raw)
    return html.unescape(parser.text())


# Render one PDF page and read it with the tesseract command-line tool.
def _ocr_page(page, lang: str, dpi: int) -> str:
    executable = shutil.which("tesseract")
    if not executable:
        raise RuntimeError("OCR needs the tesseract command-line tool")
    with tempfile.TemporaryDirectory() as tmp:
        image = Path(tmp) / "page.png"
        page.get_pixmap(dpi=dpi).save(image)
        result = subprocess.run(
            [executable, str(image), "stdout", "-l", lang],
            capture_output=True,
            text=True,
            timeout=90,
            check=False,
        )
    if result.returncode != 0:
        raise RuntimeError(f"tesseract failed: {result.stderr.strip()[:200]}")
    return result.stdout


# Read a PDF page by page, caching the extracted text by file fingerprint.
def read_pdf(
    path: Path,
    ocr: bool = False,
    ocr_lang: str = "eng",
    ocr_min_words: int = 10,
    ocr_dpi: int = 200,
) -> tuple[str, list[tuple[int, int]]]:
    try:
        import pymupdf
    except ImportError as exc:  # pragma: no cover - depends on the venv
        raise SystemExit("PDF support needs PyMuPDF; run `make venv`") from exc

    stat = path.stat()
    fingerprint = f"{path.resolve()}:{stat.st_size}:{stat.st_mtime_ns}"
    # OCR pages need a different cache entry from native text extraction.
    if ocr:
        fingerprint += f":ocr={ocr_lang}:{ocr_min_words}:{ocr_dpi}:v1"
    key = hashlib.sha1(fingerprint.encode("utf-8")).hexdigest()
    cache_file = CACHE_DIR / f"{key}.json"
    cacheable = True
    if ocr and not shutil.which("tesseract"):
        print(f"ocr skipped (tesseract not found): {path.name}", file=sys.stderr)
        ocr = False
        cacheable = False
    if cache_file.exists():
        pages = json.loads(cache_file.read_text(encoding="utf-8"))["pages"]
    else:
        pages = []
        with pymupdf.open(path) as pdf:
            for page in pdf:
                native = page.get_text("text")
                # Only image-only or nearly empty pages pay the OCR cost.
                if ocr and len(native.split()) < ocr_min_words:
                    try:
                        recognized = _ocr_page(page, ocr_lang, ocr_dpi)
                    except (OSError, RuntimeError, subprocess.SubprocessError) as exc:
                        print(f"ocr skipped ({exc}): {path.name}", file=sys.stderr)
                        cacheable = False
                        recognized = ""
                    if len(recognized.split()) > len(native.split()):
                        native = recognized
                pages.append(native)
        if cacheable:
            CACHE_DIR.mkdir(parents=True, exist_ok=True)
            cache_file.write_text(
                json.dumps({"pages": pages}, ensure_ascii=False), encoding="utf-8"
            )

    parts = []
    page_map: list[tuple[int, int]] = []
    offset = 0
    for number, page_text in enumerate(pages, start=1):
        clean = normalize_text(page_text)
        if not clean:
            continue
        if parts:
            offset += 2
        page_map.append((offset, number))
        parts.append(clean)
        offset += len(clean)
    return "\n\n".join(parts), page_map


# Read a DOCX file as plain paragraphs.
def read_docx(path: Path) -> str:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - depends on the venv
        raise SystemExit("DOCX support needs python-docx; run `make venv`") from exc
    document = docx.Document(path)
    return "\n\n".join(paragraph.text for paragraph in document.paragraphs)


# Read one supported file as text, optionally OCRing low-text PDF pages.
def read_text_file(
    path: Path,
    ocr: bool = False,
    ocr_lang: str = "eng",
    ocr_min_words: int = 10,
    ocr_dpi: int = 200,
) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return read_pdf(path, ocr, ocr_lang, ocr_min_words, ocr_dpi)[0]
    if suffix == ".docx":
        return read_docx(path)
    raw = path.read_text(encoding="utf-8", errors="replace")
    if suffix in (".html", ".htm"):
        return html_to_text(raw)
    if suffix in (".md", ".markdown"):
        return raw
    return raw


# Return the page map for a file, if its format has pages.
def page_map_for(
    path: Path,
    ocr: bool = False,
    ocr_lang: str = "eng",
    ocr_min_words: int = 10,
    ocr_dpi: int = 200,
) -> list[tuple[int, int]] | None:
    if path.suffix.lower() != ".pdf":
        return None
    return read_pdf(path, ocr, ocr_lang, ocr_min_words, ocr_dpi)[1]


# Check a relative path against the include and exclude globs.
def is_included(relative: str, include: list[str], exclude: list[str]) -> bool:
    from fnmatch import fnmatch

    # A leading **/ also matches files at the root of the source folder.
    def matches(path: str, pattern: str) -> bool:
        return fnmatch(path, pattern) or (
            pattern.startswith("**/") and fnmatch(path, pattern[3:])
        )

    if any(matches(relative, pattern) for pattern in exclude):
        return False
    return any(matches(relative, pattern) for pattern in include)


# Walk a folder and yield supported, non-excluded files in a stable order.
def iter_files(source) -> Iterator[Path]:
    root = source.path.resolve()
    for path in sorted(root.rglob("*")):
        if not path.is_file() or path.is_symlink():
            continue
        relative = path.relative_to(root).as_posix()
        if any(part.startswith(".") for part in Path(relative).parts):
            continue
        if any(part in SKIP_DIRS for part in Path(relative).parts):
            continue
        if not is_included(relative, source.include, source.exclude):
            continue
        try:
            if path.stat().st_size > source.max_file_bytes:
                print(f"skip (too large): {relative}", file=sys.stderr)
                continue
        except OSError as exc:
            print(f"skip ({exc}): {relative}", file=sys.stderr)
            continue
        yield path


# Load one private notes/papers folder as documents.
def load_file_source(source, corpus: str, limit: int | None = None) -> Iterator[SourceDocument]:
    root = source.path.resolve()
    count = 0
    for path in iter_files(source):
        if limit is not None and count >= limit:
            break
        relative = path.relative_to(root).as_posix()
        try:
            # PDFs need one pass so OCR is not repeated for the page map.
            if path.suffix.lower() == ".pdf":
                raw_text, page_map = read_pdf(
                    path,
                    ocr=source.ocr,
                    ocr_lang=source.ocr_lang,
                    ocr_min_words=source.ocr_min_words,
                    ocr_dpi=source.ocr_dpi,
                )
                text = normalize_text(raw_text)
            else:
                text = normalize_text(read_text_file(path))
                page_map = None
        except Exception as exc:  # noqa: BLE001 - one bad file must not stop the run
            print(f"skip ({exc}): {relative}", file=sys.stderr)
            continue
        if not text:
            continue
        if len(text.split()) < source.min_words:
            continue
        title = guess_title(path, text)
        metadata = {"kind": "file", "suffix": path.suffix.lower()}
        yield SourceDocument(
            id=f"{corpus}:{relative}",
            corpus=corpus,
            corpus_version="fs-1",
            title=title,
            section=f"{corpus.capitalize()} > {path.parent.relative_to(root).as_posix() or '.'}",
            path=str(path),
            sha256=content_hash(text, metadata),
            source={
                "kind": "file",
                "root": str(root),
                "relative": relative,
                "suffix": path.suffix.lower(),
                "size": path.stat().st_size,
                "mtime_ns": path.stat().st_mtime_ns,
            },
            text=text,
            page_map=page_map,
        )
        count += 1


# Turn a paper file name into a readable title.
def paper_title(filename: str) -> str:
    stem = Path(filename).stem
    stem = re.sub(r"^\d{4,6}\.\d{4,5}_", "", stem)
    return re.sub(r"[_-]+", " ", stem).strip()[:240]


# Load the lab corpus from its manifest.
def load_lab_corpus(corpus_path: Path, limit: int | None = None) -> Iterator[SourceDocument]:
    manifest = json.loads((corpus_path / "MANIFEST.json").read_text(encoding="utf-8"))
    version = manifest["corpus_version"]
    for number, entry in enumerate(manifest["documents"]):
        if limit is not None and number >= limit:
            break
        path = corpus_path / entry["path"]
        raw = path.read_text(encoding="utf-8")
        _, meta_text, body = raw.split("---\n", 2)
        meta = yaml.safe_load(meta_text) or {}
        yield SourceDocument(
            id=entry["id"],
            corpus="lab",
            corpus_version=version,
            title=entry["title"],
            section=entry["section"],
            path=entry["path"],
            sha256=entry["sha256"],
            source=meta["source"],
            text=normalize_text(body),
        )


# Load the paper library from summaries and processed PDFs.
def load_papers(source, limit: int | None = None) -> Iterator[SourceDocument]:
    root = source.path.resolve()
    pdf_dir = root / (source.pdf_dir or "processed-pdfs")
    summaries_dir = root / (source.summaries_dir or "analysis_results_per_file")
    summaries: dict[str, dict] = {}

    # Read the summaries first; they carry the tags and the clean abstract.
    if summaries_dir.exists():
        for path in sorted(summaries_dir.glob("*.json")):
            try:
                data = json.loads(path.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError) as exc:
                print(f"skip summary ({exc}): {path.name}", file=sys.stderr)
                continue
            filename = data.get("filename")
            if not isinstance(filename, str) or not filename or "?skip=" in filename:
                continue
            summaries[filename] = data

    filenames = list(summaries)
    for path in sorted(pdf_dir.glob("*.pdf")):
        if path.name not in summaries:
            filenames.append(path.name)

    count = 0
    for filename in filenames:
        if limit is not None and count >= limit:
            break
        data = summaries.get(filename, {})
        summary_value = data.get("summary") or ""
        summary = summary_value.strip() if isinstance(summary_value, str) else str(summary_value)
        tags = data.get("tags") or []
        if isinstance(tags, str):
            tags = [tags]
        tags = [
            str(tag).strip().replace("\x00", "")
            for tag in tags
            if str(tag).strip().replace("\x00", "")
        ]
        pdf_path = pdf_dir / filename
        pdf_text = ""
        page_map = None
        if pdf_path.exists():
            try:
                pdf_text, page_map = read_pdf(
                    pdf_path,
                    ocr=source.ocr,
                    ocr_lang=source.ocr_lang,
                    ocr_min_words=source.ocr_min_words,
                    ocr_dpi=source.ocr_dpi,
                )
            except Exception as exc:  # noqa: BLE001 - one bad PDF must not stop the run
                print(f"skip pdf ({exc}): {pdf_path.name}", file=sys.stderr)

        text_parts = [paper_title(filename)]
        if summary:
            text_parts.append(f"Summary: {summary}")
        if tags:
            text_parts.append("Topics: " + ", ".join(str(tag) for tag in tags))
        if pdf_text:
            text_parts.append(pdf_text)
        text = normalize_text("\n\n".join(text_parts))
        if len(text.split()) < 30:
            continue

        metadata = {"tags": tags, "summary": summary, "pdf": pdf_path.exists()}
        source_info = {
            "kind": "paper",
            "pdf": str(pdf_path) if pdf_path.exists() else None,
            "summary": str(summaries_dir / f"{Path(filename).stem}.json")
            if data
            else None,
            "tags": tags,
        }
        yield SourceDocument(
            id=f"papers:{filename}",
            corpus="papers",
            corpus_version="papers-1",
            title=paper_title(filename),
            section="Papers > " + (str(tags[0]) if tags else "Unfiled"),
            path=str(pdf_path) if pdf_path.exists() else str(summaries_dir / f"{Path(filename).stem}.json"),
            sha256=content_hash(text, metadata),
            source=source_info,
            text=text,
            page_map=page_map,
        )
        count += 1


# Resolve a corpus name to its document iterator.
def iter_documents(
    settings, corpus: str, limit: int | None = None, ocr: bool | None = None
) -> Iterator[SourceDocument]:
    if corpus == "lab":
        corpus_path = Path(__file__).resolve().parents[1] / "corpus"
        return load_lab_corpus(corpus_path, limit=limit)
    source = settings.source(corpus)
    # A CLI override applies to this run without editing the private config.
    if ocr is not None:
        source = replace(source, ocr=ocr)
    if source.kind == "papers":
        return load_papers(source, limit=limit)
    return load_file_source(source, corpus, limit=limit)
